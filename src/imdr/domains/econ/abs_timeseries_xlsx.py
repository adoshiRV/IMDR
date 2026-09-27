"""Shared parser for ABS "Time Series Workbook" XLSX releases.

Distinct from ``abs_sdmx.py`` (the SDMX REST API): a handful of ABS tables
are only published as standalone time-series workbooks on the publication
page rather than through the SDMX data API (e.g. ``6291004.xlsx`` — Labour
Force, Australia, Detailed, Table 04, Employed persons by Industry
division). Those workbooks all share one layout:

  - An ``Index`` sheet: a header row (columns ``Data Item Description``,
    ``Series Type``, ``Series ID``, ``Series Start``, ``Series End``,
    ``No. Obs.``, ``Unit``, ``Data Type``, ``Freq.``, ``Collection Month``)
    followed by one row per series mapping ``Series ID`` -> its description
    + series type (``Original`` / ``Seasonally Adjusted`` / ``Trend``).
  - One or more ``Data1``, ``Data2``, ... sheets: a header row with
    ``Series ID`` in column A and the series ID codes across the remaining
    columns, followed by data rows (date in column A, one value per series
    column). Verified live against ``6291004.xlsx`` 2026-07-23.

The publication landing page ("latest-release") embeds the current
quarterly/monthly release folder (e.g. ``mar-2026``) as a relative href
ending in the table filename; that folder segment changes every release,
so callers should resolve it via ``resolve_latest_xlsx_url`` rather than
hardcode a URL.
"""
from __future__ import annotations

import datetime
import re
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

import httpx

UTC = datetime.timezone.utc
_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
_ABS_BASE = "https://www.abs.gov.au"


@dataclass(frozen=True)
class ABSXlsxSeries:
    """One series parsed out of an ABS time-series workbook."""

    series_id: str
    description: str
    series_type: str  # "Original" | "Seasonally Adjusted" | "Trend"
    obs: list[tuple[datetime.date, float | None]] = field(default_factory=list)


def resolve_latest_xlsx_url(landing_url: str, table_filename: str, *, timeout: float = 30.0) -> str:
    """Resolve the current release URL for an ABS time-series table.

    ABS embeds the release folder (which changes every release, e.g.
    ``mar-2026``) as a relative href on the ``latest-release`` landing page.
    Regexing it out here means callers never hardcode a stale folder name.
    """
    r = httpx.get(landing_url, headers={"User-Agent": _UA}, timeout=timeout, follow_redirects=True)
    r.raise_for_status()
    pattern = re.compile(r'href="([^"]*' + re.escape(table_filename) + r')"')
    m = pattern.search(r.text)
    if not m:
        raise ValueError(
            f"could not find an href ending {table_filename!r} on {landing_url}"
        )
    href = m.group(1)
    return href if href.startswith("http") else _ABS_BASE + href


def download_workbook(url: str, *, timeout: float = 60.0) -> bytes:
    r = httpx.get(url, headers={"User-Agent": _UA}, timeout=timeout, follow_redirects=True)
    r.raise_for_status()
    return r.content


def _norm(v: object) -> str:
    return str(v).strip() if v is not None else ""


def _find_header_row(rows: list[list], marker: str, *, col_idx: int = 0, max_scan: int = 20) -> int:
    for i, row in enumerate(rows[:max_scan]):
        if col_idx < len(row) and _norm(row[col_idx]).lower() == marker.lower():
            return i
    raise RuntimeError(f"could not locate header row (looking for {marker!r} in column {col_idx})")


def parse_index_sheet(ws) -> dict[str, tuple[str, str]]:
    """Return ``{series_id: (description, series_type)}`` from the Index sheet."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i = _find_header_row(rows, "Data Item Description", col_idx=0)
    header = [_norm(c) for c in rows[hdr_i]]
    try:
        i_desc = header.index("Data Item Description")
        i_type = header.index("Series Type")
        i_id = header.index("Series ID")
    except ValueError as exc:
        raise RuntimeError(f"Index sheet header missing expected column: {header!r}") from exc

    out: dict[str, tuple[str, str]] = {}
    for row in rows[hdr_i + 1:]:
        if i_id >= len(row) or not row[i_id]:
            continue
        sid = _norm(row[i_id])
        desc = _norm(row[i_desc]) if i_desc < len(row) else ""
        stype = _norm(row[i_type]) if i_type < len(row) else ""
        out[sid] = (desc, stype)
    return out


def parse_data_sheet(ws) -> dict[str, list[tuple[datetime.date, float | None]]]:
    """Return ``{series_id: [(date, value), ...]}`` from one Data sheet."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i = _find_header_row(rows, "Series ID", col_idx=0)
    header = rows[hdr_i]
    col_series: dict[int, str] = {}
    for j, cell in enumerate(header):
        if j == 0:
            continue
        sid = _norm(cell)
        if sid:
            col_series[j] = sid

    out: dict[str, list[tuple[datetime.date, float | None]]] = {sid: [] for sid in col_series.values()}
    for row in rows[hdr_i + 1:]:
        if not row or row[0] is None:
            continue
        raw_date = row[0]
        if isinstance(raw_date, datetime.datetime):
            d = raw_date.date()
        elif isinstance(raw_date, datetime.date):
            d = raw_date
        else:
            continue
        for j, sid in col_series.items():
            # Inline float coercion (rather than reusing aofm_xlsx.coerce_float):
            # ABS time-series-workbook data cells are always plain numeric floats
            # or blanks with data_only=True — no 'TBA'/'na'/comma-string cases that
            # the AOFM sheets carry, so the blank-or-float logic here suffices.
            v = row[j] if j < len(row) else None
            if v is None or v == "":
                fv: float | None = None
            else:
                try:
                    fv = float(v)
                except (TypeError, ValueError):
                    fv = None
            out[sid].append((d, fv))
    return out


def parse_workbook(data: bytes | Path) -> dict[str, ABSXlsxSeries]:
    """Parse an ABS time-series workbook into ``{series_id: ABSXlsxSeries}``.

    ``data`` may be raw bytes (as returned by ``download_workbook``) or a
    filesystem path (for tests / cached fixtures).
    """
    import openpyxl

    src = BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    try:
        if "Index" not in wb.sheetnames:
            raise RuntimeError(f"workbook missing Index sheet; got {wb.sheetnames!r}")
        index_map = parse_index_sheet(wb["Index"])

        obs_map: dict[str, list[tuple[datetime.date, float | None]]] = {}
        for name in wb.sheetnames:
            if not name.lower().startswith("data"):
                continue
            for sid, obs in parse_data_sheet(wb[name]).items():
                obs_map.setdefault(sid, []).extend(obs)

        out: dict[str, ABSXlsxSeries] = {}
        for sid, (desc, stype) in index_map.items():
            obs = sorted(obs_map.get(sid, []), key=lambda t: t[0])
            out[sid] = ABSXlsxSeries(series_id=sid, description=desc, series_type=stype, obs=obs)
        return out
    finally:
        wb.close()
