"""Unit tests for the ABS time-series XLSX parser + the LF industry-division
fetcher (scripts.econ.au.abs.abs_lf_industry).

No network: workbooks are built in-memory with openpyxl, mirroring the real
6291004.xlsx layout verified live 2026-07-23 (Index sheet header at row 10 /
Data sheet ``Series ID`` header at row 10, values below).
"""
from __future__ import annotations

import datetime
from io import BytesIO

import openpyxl
import pytest

from imdr.domains.econ.abs_timeseries_xlsx import (
    parse_data_sheet,
    parse_index_sheet,
    parse_workbook,
    resolve_latest_xlsx_url,
)
from imdr.domains.econ.schema import VALID_CATEGORIES, VALID_FREQUENCIES
from scripts.econ.au.abs import abs_lf_industry


def _build_workbook(*, divisions: list[str], dates: list[datetime.date]) -> bytes:
    """Build a synthetic ABS time-series workbook: Index + Data1 sheets,
    ``divisions`` each carrying Trend/Seasonally Adjusted/Original series,
    with a linearly increasing value per (division, series_type) so QoQ
    deltas are deterministic and non-zero.
    """
    wb = openpyxl.Workbook()
    index_ws = wb.active
    index_ws.title = "Index"
    for _ in range(9):
        index_ws.append([None] * 12)
    index_ws.append([
        "Data Item Description", None, None, "Series Type", "Series ID",
        "Series Start", "Series End", "No. Obs.", "Unit", "Data Type",
        "Freq.", "Collection Month",
    ])

    data_ws = wb.create_sheet("Data1")
    sid_row: list = ["Series ID"]

    counter = 0
    for div in divisions:
        desc = "Employed total ;" if div == "Employed total ;" else f"{div} ;  Employed total ;"
        for stype in ("Trend", "Seasonally Adjusted", "Original"):
            counter += 1
            sid = f"S{counter:04d}"
            index_ws.append([
                desc, None, None, stype, sid,
                dates[0], dates[-1], len(dates), "000", "STOCK", "Quarter", 2,
            ])
            sid_row.append(sid)

    # 9 blank rows then the header => "Series ID" lands on row 10, matching the
    # real 6291004.xlsx Data-sheet layout (verified live 2026-07-23).
    for _ in range(9):
        data_ws.append([None] * len(sid_row))
    data_ws.append(sid_row)

    for i, d in enumerate(dates):
        row = [datetime.datetime(d.year, d.month, d.day)]
        for div in divisions:
            for stype in ("Trend", "Seasonally Adjusted", "Original"):
                row.append(100.0 + 10.0 * i + hash((div, stype)) % 3)
        data_ws.append(row)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParseIndexSheet:
    def test_maps_series_id_to_description_and_type(self) -> None:
        data = _build_workbook(
            divisions=["Construction", "Mining"],
            dates=[datetime.date(2025, 5, 1), datetime.date(2025, 8, 1)],
        )
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        index_map = parse_index_sheet(wb["Index"])
        assert len(index_map) == 6
        descs_types = set(index_map.values())
        assert ("Construction ;  Employed total ;", "Trend") in descs_types
        assert ("Mining ;  Employed total ;", "Original") in descs_types


class TestParseDataSheet:
    def test_reads_dates_and_values_per_series(self) -> None:
        dates = [datetime.date(2025, 5, 1), datetime.date(2025, 8, 1), datetime.date(2025, 11, 1)]
        data = _build_workbook(divisions=["Construction"], dates=dates)
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        obs_map = parse_data_sheet(wb["Data1"])
        assert len(obs_map) == 3
        for sid, obs in obs_map.items():
            assert [d for d, _ in obs] == dates
            assert all(v is not None for _, v in obs)


class TestParseWorkbook:
    def test_joins_index_and_data(self) -> None:
        dates = [datetime.date(2025, 5, 1), datetime.date(2025, 8, 1), datetime.date(2025, 11, 1)]
        data = _build_workbook(divisions=["Construction", "Employed total ;"], dates=dates)
        parsed = parse_workbook(data)
        assert len(parsed) == 6
        for series in parsed.values():
            assert series.series_type in {"Trend", "Seasonally Adjusted", "Original"}
            assert len(series.obs) == 3
            assert [d for d, _ in series.obs] == dates

    def test_accepts_path(self, tmp_path) -> None:
        dates = [datetime.date(2025, 5, 1), datetime.date(2025, 8, 1)]
        data = _build_workbook(divisions=["Mining"], dates=dates)
        p = tmp_path / "wb.xlsx"
        p.write_bytes(data)
        parsed = parse_workbook(p)
        assert len(parsed) == 3


class TestResolveLatestXlsxUrl:
    def test_resolves_relative_href(self, monkeypatch: pytest.MonkeyPatch) -> None:
        html = '<html><a href="/statistics/foo/mar-2026/6291004.xlsx">Table 04</a></html>'

        class _Resp:
            status_code = 200
            text = html

            def raise_for_status(self) -> None:
                return None

        def _fake_get(url, headers=None, timeout=None, follow_redirects=None):
            return _Resp()

        monkeypatch.setattr("imdr.domains.econ.abs_timeseries_xlsx.httpx.get", _fake_get)
        url = resolve_latest_xlsx_url("https://www.abs.gov.au/landing", "6291004.xlsx")
        assert url == "https://www.abs.gov.au/statistics/foo/mar-2026/6291004.xlsx"

    def test_raises_when_not_found(self, monkeypatch: pytest.MonkeyPatch) -> None:
        class _Resp:
            status_code = 200
            text = "<html>nothing here</html>"

            def raise_for_status(self) -> None:
                return None

        monkeypatch.setattr(
            "imdr.domains.econ.abs_timeseries_xlsx.httpx.get",
            lambda *a, **kw: _Resp(),
        )
        with pytest.raises(ValueError, match="6291004.xlsx"):
            resolve_latest_xlsx_url("https://www.abs.gov.au/landing", "6291004.xlsx")


class TestDivisionKey:
    @pytest.mark.parametrize("description,expected", [
        ("Education and Training ;  Employed total ;", "Education and Training"),
        ("Health Care and Social Assistance ;  Employed total ;", "Health Care and Social Assistance"),
        ("Employed total ;", abs_lf_industry._TOTAL_KEY),
    ])
    def test_extracts_division_or_total(self, description: str, expected: str) -> None:
        assert abs_lf_industry._division_key(description) == expected

    def test_unrecognised_returns_first_segment(self) -> None:
        assert abs_lf_industry._division_key("Some Unknown Bucket ;  Employed total ;") == "Some Unknown Bucket"


class TestDivisionsMapping:
    def test_19_divisions_plus_total(self) -> None:
        assert len(abs_lf_industry._DIVISIONS) == 20
        assert abs_lf_industry._TOTAL_KEY in abs_lf_industry._DIVISIONS

    def test_suffixes_unique(self) -> None:
        suffixes = list(abs_lf_industry._DIVISIONS.values())
        assert len(suffixes) == len(set(suffixes)), "duplicate imdr_code suffixes"

    def test_suffixes_are_upper_snake(self) -> None:
        for suffix in abs_lf_industry._DIVISIONS.values():
            assert suffix == suffix.upper()
            assert " " not in suffix


class TestRunFetch:
    def test_builds_indicators_and_observations(self, monkeypatch: pytest.MonkeyPatch) -> None:
        dates = [datetime.date(2025, 8, 1), datetime.date(2025, 11, 1), datetime.date(2026, 2, 1)]
        data = _build_workbook(
            divisions=["Construction", "Education and Training", "Employed total ;"],
            dates=dates,
        )
        monkeypatch.setattr(abs_lf_industry, "resolve_latest_xlsx_url", lambda *a, **kw: "https://fake/6291004.xlsx")
        monkeypatch.setattr(abs_lf_industry, "download_workbook", lambda url: data)

        indicators, observations = abs_lf_industry.run_fetch(None, None)

        assert len(indicators) == 9  # 3 divisions x 3 series types
        codes = {i.imdr_code for i in indicators}
        assert "ABS.LF.EMPLOYED_IND_CONSTRUCTION_SA.AU" in codes
        assert "ABS.LF.EMPLOYED_IND_EDUCATION_TRAINING_TREND.AU" in codes
        assert "ABS.LF.EMPLOYED_IND_ALL_INDUSTRIES_ORIG.AU" in codes

        for ind in indicators:
            assert ind.category in VALID_CATEGORIES
            assert ind.frequency in VALID_FREQUENCIES
            assert ind.unit == "persons"
            assert ind.country_iso == "AU"
            assert ind.imdr_code.endswith(".AU")

        sa_flags = {i.imdr_code: i.is_seasonally_adjusted for i in indicators}
        assert sa_flags["ABS.LF.EMPLOYED_IND_CONSTRUCTION_SA.AU"] is True
        assert sa_flags["ABS.LF.EMPLOYED_IND_CONSTRUCTION_TREND.AU"] is False
        assert sa_flags["ABS.LF.EMPLOYED_IND_CONSTRUCTION_ORIG.AU"] is False

        assert len(observations) == 9 * len(dates)

    def test_since_until_filter_observations(self, monkeypatch: pytest.MonkeyPatch) -> None:
        dates = [datetime.date(2025, 8, 1), datetime.date(2025, 11, 1), datetime.date(2026, 2, 1)]
        data = _build_workbook(divisions=["Construction"], dates=dates)
        monkeypatch.setattr(abs_lf_industry, "resolve_latest_xlsx_url", lambda *a, **kw: "https://fake/6291004.xlsx")
        monkeypatch.setattr(abs_lf_industry, "download_workbook", lambda url: data)

        _, observations = abs_lf_industry.run_fetch("2025-11-01", None)
        assert all(o.obs_date >= datetime.date(2025, 11, 1) for o in observations)
        assert len(observations) == 3 * 2  # 3 series types x 2 remaining quarters
